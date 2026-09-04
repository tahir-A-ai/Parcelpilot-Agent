"""
ETL Script: ParcelPilot_Assessment_Data.xlsx → SQLite

Reads the three data sheets (accounts, orders, tickets) and upserts every
row into the SQLite database. The operation is fully idempotent — re-running
this script produces no duplicates and updates any changed values in place.

Usage (from Backend/ directory):
    python scripts/ingest_sqlite.py

Design decisions:
- `session.merge()` provides upsert semantics: INSERT if PK absent, UPDATE if present.
- Each table is committed independently so a failure in one sheet does not
  roll back successful inserts in earlier sheets.
- Type coercion helpers normalise the inconsistent types openpyxl may return
  (datetime objects, string timestamps, string booleans, etc.).
"""

from __future__ import annotations

import asyncio
import logging
import sys
from datetime import datetime
from pathlib import Path

# Ensure Backend/ is on sys.path so `app.*` imports resolve correctly
# regardless of which directory the script is invoked from.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.model.models import Account, Order, Ticket
from app.db.session import AsyncSessionLocal, init_db

logging.basicConfig(level=logging.INFO, format="%(levelname)s  %(message)s")
logger = logging.getLogger(__name__)

EXCEL_PATH = (
    Path(__file__).resolve().parent.parent / "data" / "raw" / "ParcelPilot_Assessment_Data.xlsx"
)


# --------------------------------------------------------------------------- #
# Type-coercion helpers                                                        #
# --------------------------------------------------------------------------- #

def _str(value: object) -> str | None:
    """Return a stripped string or None for empty / null-like values."""
    if value is None:
        return None
    s = str(value).strip()
    return s if s and s.lower() not in ("none", "null") else None


def _bool(value: object) -> bool:
    """Coerce Excel boolean, integer, or string to Python bool."""
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return bool(value)
    if value is None:
        return False
    return str(value).strip().lower() in ("true", "1", "yes")


def _float(value: object) -> float:
    """Coerce value to float, defaulting to 0.0 on failure."""
    if value is None:
        return 0.0
    try:
        return float(value)
    except (ValueError, TypeError):
        return 0.0


def _dt(value: object) -> datetime | None:
    """
    Coerce a cell value to datetime.

    Handles:
    - Native datetime objects returned by openpyxl for date-formatted cells.
    - ISO-like string timestamps stored as plain text ("2026-08-16 09:00").
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    s = str(value).strip()
    if not s or s.lower() in ("none", "null"):
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        logger.warning("Unrecognised datetime value, skipping: %r", s)
        return None


def _sheet_to_dicts(ws) -> list[dict]:
    """
    Convert a worksheet to a list of row dicts keyed by the header row.
    Completely empty rows are skipped.
    """
    headers = [cell.value for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(dict(zip(headers, row)))
    return rows


# --------------------------------------------------------------------------- #
# Per-sheet ingest functions                                                   #
# --------------------------------------------------------------------------- #

async def _ingest_accounts(session: AsyncSession, wb: openpyxl.Workbook) -> int:
    rows = _sheet_to_dicts(wb["accounts"])
    for r in rows:
        await session.merge(Account(
            account_id=_str(r["account_id"]),
            account_name=_str(r["account_name"]),
            plan=_str(r["plan"]),
            status=_str(r["status"]),
            csm=_str(r.get("csm")),
            contract_file=_str(r.get("contract_file")),
            premium_support=_bool(r.get("premium_support", False)),
            notes=_str(r.get("notes")),
        ))
    return len(rows)


async def _ingest_orders(session: AsyncSession, wb: openpyxl.Workbook) -> int:
    rows = _sheet_to_dicts(wb["orders"])
    for r in rows:
        await session.merge(Order(
            order_id=_str(r["order_id"]),
            account_id=_str(r["account_id"]),
            carrier=_str(r["carrier"]),
            status=_str(r["status"]),
            booked_at=_dt(r["booked_at"]),
            pickup_window_start=_dt(r["pickup_window_start"]),
            pickup_window_end=_dt(r["pickup_window_end"]),
            pickup_actual_at=_dt(r.get("pickup_actual_at")),
            shipment_fee_inr=_float(r["shipment_fee_inr"]),
            carrier_fault=_bool(r.get("carrier_fault", False)),
            customer_fault=_bool(r.get("customer_fault", False)),
            cancellation_requested_at=_dt(r.get("cancellation_requested_at")),
            notes=_str(r.get("notes")),
        ))
    return len(rows)


async def _ingest_tickets(session: AsyncSession, wb: openpyxl.Workbook) -> int:
    rows = _sheet_to_dicts(wb["tickets"])
    for r in rows:
        await session.merge(Ticket(
            ticket_id=_str(r["ticket_id"]),
            account_id=_str(r["account_id"]),
            created_at=_dt(r["created_at"]),
            status=_str(r["status"]),
            subject=_str(r["subject"]),
            description=_str(r.get("description")) or "",
            channel=_str(r["channel"]),
            assigned_to=_str(r.get("assigned_to")),
            last_customer_message_at=_dt(r["last_customer_message_at"]),
            historical_resolution=_str(r.get("historical_resolution")),
        ))
    return len(rows)


# --------------------------------------------------------------------------- #
# Entry point                                                                  #
# --------------------------------------------------------------------------- #

async def main() -> None:
    if not EXCEL_PATH.exists():
        logger.error("Excel file not found: %s", EXCEL_PATH)
        sys.exit(1)

    logger.info("Loading workbook: %s", EXCEL_PATH)
    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)

    logger.info("Ensuring database tables exist...")
    await init_db()

    ingest_tasks = [
        ("accounts", _ingest_accounts),
        ("orders",   _ingest_orders),
        ("tickets",  _ingest_tickets),
    ]

    async with AsyncSessionLocal() as session:
        for table_name, fn in ingest_tasks:
            try:
                count = await fn(session, wb)
                await session.commit()
                logger.info("  [%-10s] %d rows upserted.", table_name, count)
            except Exception as exc:
                await session.rollback()
                logger.error("  [%-10s] FAILED — rolling back: %s", table_name, exc)
                raise

    logger.info("SQLite ingestion complete.")


if __name__ == "__main__":
    asyncio.run(main())
